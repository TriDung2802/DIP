from datetime import datetime
from pathlib import Path
import cv2
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, Alignment, PatternFill

THU_MUC_LOG = Path(__file__).resolve().parent / "logs"
THU_MUC_ANH = THU_MUC_LOG / "anh"
FILE_EXCEL = THU_MUC_LOG / "vi_pham.xlsx"
CAC_COT = ["thoi_gian", "camera", "loai_vi_pham","do_tin_cay", "file_anh", "anh_vi_pham"]

class Logger:
    def __init__(self):
        THU_MUC_ANH.mkdir(parents=True, exist_ok=True)
        if not FILE_EXCEL.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = "Vi phạm"
            ws.append(CAC_COT)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1B2A4A")
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

            for cot, rong in zip("ABCDEF", [22, 14, 24, 15, 38, 22]):
                ws.column_dimensions[cot].width = rong
            ws.row_dimensions[1].height = 25
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = "A1:F1"
            wb.save(FILE_EXCEL)

    def ghi(self, su_kien):
        try:
            moc = datetime.strptime(
                su_kien["thoi_gian"], "%d/%m/%Y %H:%M:%S"
            )
            ten_anh = f"vi_pham_{moc:%Y%m%d_%H%M%S_%f}.jpg"
            duong_dan_anh = THU_MUC_ANH / ten_anh

            anh = su_kien.get("anh")
            if anh is None or anh.size == 0:
                raise ValueError("Không có ảnh để lưu")

            if not cv2.imwrite(str(duong_dan_anh), anh):
                raise RuntimeError("Không lưu được ảnh vi phạm")

            wb = load_workbook(FILE_EXCEL)
            ws = wb.active
            dong = ws.max_row + 1

            gia_tri = [
                su_kien["thoi_gian"],
                su_kien["camera"],
                su_kien["loai"],
                float(su_kien["tin_cay"]),
                ten_anh,
            ]

            for cot, value in enumerate(gia_tri, start=1):
                o = ws.cell(dong, cot, value)
                o.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

            ws.cell(dong, 4).number_format = "0.00%"
            img = ExcelImage(str(duong_dan_anh))
            ti_le = min(150 / img.width, 90 / img.height, 1)
            img.width = int(img.width * ti_le)
            img.height = int(img.height * ti_le)
            ws.add_image(img, f"F{dong}")
            ws.row_dimensions[dong].height = 70
            wb.save(FILE_EXCEL)

            print(f"[Logger] Đã lưu ảnh: {duong_dan_anh}")
            print(f"[Logger] Đã cập nhật Excel: {FILE_EXCEL}")

        except Exception as loi:
            print(f"[Logger] Không ghi được log: {loi}")
